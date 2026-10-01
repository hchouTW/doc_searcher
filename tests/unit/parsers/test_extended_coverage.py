"""Office fixtures catch omitted comments, grouped shapes and duplicate table cells."""

from doc_searcher.parsers.xlsx_parser import XlsxParser
from doc_searcher.parsers.pptx_parser import PptxParser


def test_comments_on_empty_hidden_cells_are_searchable(tmp_path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    workbook = Workbook()
    workbook.active["A1"] = "visible"
    sheet = workbook.create_sheet("隱藏表")
    sheet.sheet_state = "hidden"
    sheet.row_dimensions[42].hidden = True
    sheet["C42"].comment = Comment("會議 comment only", "author")
    path = tmp_path / "comments.xlsx"
    workbook.save(path)
    result = XlsxParser().parse(str(path))
    assert result.full_text.count("會議") == 1
    source = next(
        span["source"]
        for segment in result.segments
        for span in segment.sources
        if span["source"]["kind"] == "comment"
    )
    assert source["location"] == "隱藏表!C42"


def test_nested_group_shapes_and_merged_cells(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "會議 grouped"
    table = slide.shapes.add_table(1, 2, 0, 0, Inches(2), Inches(1)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "merged content"
    path = tmp_path / "groups.pptx"
    presentation.save(path)
    result = PptxParser().parse(str(path))
    assert result.full_text.count("會議") == 1
    assert result.full_text.count("merged content") == 1
    assert any(span["source"]["kind"] == "shape" for span in result.segments[0].sources)


def test_nonrendered_master_shapes_are_excluded(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1))
    shape.text = "HiddenMasterNeedle"
    presentation.slide_master.shapes._spTree.append(shape._element)
    slide._element.set("showMasterSp", "0")
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "visible"
    path = tmp_path / "hidden-master.pptx"
    presentation.save(path)
    result = PptxParser().parse(str(path))
    assert "visible" in result.full_text and "HiddenMasterNeedle" not in result.full_text


def test_xlsx_embedded_object_reports_partial_coverage(tmp_path):
    import zipfile
    from lxml import etree
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active["A1"] = "readable"
    path = tmp_path / "embedded.xlsx"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    relations = etree.Element("{" + namespace + "}Relationships", nsmap={None: namespace})
    etree.SubElement(
        relations,
        "{" + namespace + "}Relationship",
        Id="ole1",
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
        Target="../embeddings/object.bin",
    )
    contents["xl/worksheets/_rels/sheet1.xml.rels"] = etree.tostring(relations)
    contents["xl/embeddings/object.bin"] = b"embedded payload"
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in contents.items():
            archive.writestr(name, content)
    result = XlsxParser().parse(str(path))
    assert result.status.value == "partial"
    assert result.full_text == "readable"
    assert any(w["code"] == "unsupported_embedded_object" for w in result.warnings)
