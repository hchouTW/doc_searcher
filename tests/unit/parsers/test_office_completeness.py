"""Synthetic Office source locations and partial PDF extraction regressions."""
import zipfile
from unittest.mock import MagicMock

from doc_searcher.parsers.docx_parser import DocxParser
from doc_searcher.parsers.xlsx_parser import XlsxParser
from doc_searcher.parsers.pdf_parser import PdfParser


def test_docx_order_shared_parts_and_text_boxes(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    doc = Document()
    doc.add_paragraph("before")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "same"
    table.cell(0, 1).text = "same"
    doc.add_paragraph("after")
    box = OxmlElement("w:txbxContent")
    paragraph = OxmlElement("w:p")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "text-box-only"
    run.append(text)
    paragraph.append(run)
    box.append(paragraph)
    doc.paragraphs[0]._p.append(box)
    doc.sections[0].header.paragraphs[0].text = "header-only"
    doc.sections[0].footer.paragraphs[0].text = "footer-only"
    doc.add_section()
    path = tmp_path / "parts.docx"
    doc.save(path)
    extracted = DocxParser().parse(str(path))
    text = extracted.full_text
    assert text.index("before") < text.index("same") < text.index("after")
    assert text.count("same") == 2
    for term in ("header-only", "footer-only", "text-box-only"):
        assert text.count(term) == 1
    sources = [span["source"] for segment in extracted.segments for span in segment.sources]
    assert {s["kind"] for s in sources} >= {"header", "footer", "text_box", "paragraph", "table_cell"}
    assert all("part" in s for s in sources)


def test_xlsx_formulas_caches_cells_and_hidden_sheets(tmp_path):
    import openpyxl
    from lxml import etree
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in range(1, 9):
        sheet.cell(row, 2, "會議")
    sheet["C1"] = '="cache-only"'
    sheet["C2"] = '="uncached-only"'
    sheet["D1"] = "merged-only"
    sheet.merge_cells("D1:E1")
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "hidden-only"
    path = tmp_path / "formulas.xlsx"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        contents = {n: archive.read(n) for n in archive.namelist()}
    xml = etree.fromstring(contents["xl/worksheets/sheet1.xml"])
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cell = xml.xpath('//s:c[@r="C1"]', namespaces=ns)[0]
    cell.set("t", "str")
    cell.find("s:v", namespaces=ns).text = "cached-result"
    contents["xl/worksheets/sheet1.xml"] = etree.tostring(xml)
    with zipfile.ZipFile(path, "w") as archive:
        for name, value in contents.items():
            archive.writestr(name, value)
    result = XlsxParser().parse(str(path))
    assert all(x in result.full_text for x in ("cache-only", "cached-result", "uncached-only", "hidden-only"))
    assert result.full_text.count("會議") == 8
    assert result.full_text.count("merged-only") == 1
    assert result.warnings and any("C2" in str(x) for x in result.warnings)
    spans = [x for segment in result.segments for x in segment.sources]
    assert {x["source"]["kind"] for x in spans} >= {"value", "formula"}
    assert {x["source"]["location"] for x in spans} >= {"Sheet1!B1", "Sheet1!B8", "Sheet1!C2"}
    for segment in result.segments:
        for span in segment.sources:
            assert segment.text[span["start"]:span["end"]]


def test_pdf_partial_page_failure_is_recorded(tmp_path, monkeypatch):
    import pymupdf
    doc = MagicMock()
    doc.needs_pass = False
    doc.__len__.return_value = 3
    page = MagicMock()
    page.get_text.return_value = "readable-page"
    doc.load_page.side_effect = [page, RuntimeError("broken page"), page]
    monkeypatch.setattr(pymupdf, "open", lambda _: doc)
    result = PdfParser().parse(str(tmp_path / "partial.pdf"))
    assert result.status.value == "partial"
    assert len(result.segments) == 2
    assert result.omitted_locations == ["page 2"]
    assert any("2" in str(x) for x in result.warnings)
    doc.close.assert_called_once()


def test_formula_cached_text_is_not_counted_twice(tmp_path):
    import openpyxl
    from lxml import etree
    wb = openpyxl.Workbook()
    wb.active["A1"] = '="會議"'
    path = tmp_path / "same-source.xlsx"
    wb.save(path)
    with zipfile.ZipFile(path) as archive:
        contents = {n: archive.read(n) for n in archive.namelist()}
    root = etree.fromstring(contents["xl/worksheets/sheet1.xml"])
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    cell = root.xpath('//s:c[@r="A1"]', namespaces=namespace)[0]
    cell.set("t", "str")
    cell.find("s:v", namespaces=namespace).text = "會議"
    contents["xl/worksheets/sheet1.xml"] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in contents.items():
            archive.writestr(name, data)
    result = XlsxParser().parse(str(path))
    assert result.full_text.count("會議") == 1
    assert result.segments[0].sources[0]["source"]["cached_value"] == "會議"


def test_docx_alternate_textbox_has_one_effective_representation(tmp_path):
    from docx import Document
    from lxml import etree
    doc = Document()
    paragraph = doc.add_paragraph("body")
    paragraph._p.append(etree.fromstring('''
      <mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
       xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
       xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <mc:Choice Requires="wps"><w:txbxContent><w:p><w:r><w:t>UniqueTextBox</w:t></w:r></w:p></w:txbxContent></mc:Choice>
        <mc:Fallback><w:txbxContent><w:p><w:r><w:t>UniqueTextBox</w:t></w:r></w:p></w:txbxContent></mc:Fallback>
      </mc:AlternateContent>'''))
    # Unknown required namespace must select the fallback, not silently extract both.
    paragraph._p.append(etree.fromstring('''
      <mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
       xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:unknown="urn:unsupported">
        <mc:Choice Requires="unknown"><w:p><w:r><w:t>unsupported-choice</w:t></w:r></w:p></mc:Choice>
        <mc:Fallback><w:p><w:r><w:t>fallback-text</w:t></w:r></w:p></mc:Fallback>
      </mc:AlternateContent>'''))
    path = tmp_path / "alternate.docx"
    doc.save(path)
    result = DocxParser().parse(str(path))
    assert result.full_text.count("UniqueTextBox") == 1
    assert "fallback-text" in result.full_text
    assert "unsupported-choice" not in result.full_text


def test_xlsx_array_formula_source_and_range(tmp_path):
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    wb = openpyxl.Workbook()
    wb.active["A1"] = ArrayFormula(ref="A1:A3", text='=IF(B1:B3="NeedleFormula",1,0)')
    path = tmp_path / "array.xlsx"
    wb.save(path)
    result = XlsxParser().parse(str(path))
    assert "NeedleFormula" in result.full_text
    assert "object at" not in result.full_text
    source = result.segments[0].sources[0]["source"]
    assert source["range"] == "A1:A3"
    assert source["kind"] == "formula"
    assert any(w["code"] == "missing_cached_result" for w in result.warnings)


def test_xlsx_nontext_formula_reports_omission(tmp_path):
    import openpyxl
    from openpyxl.worksheet.formula import DataTableFormula
    wb = openpyxl.Workbook()
    wb.active["A1"] = DataTableFormula(ref="A1:A3", r1="B1")
    path = tmp_path / "data-table.xlsx"
    wb.save(path)
    result = XlsxParser().parse(str(path))
    assert "object at" not in result.full_text
    assert any(w["code"] == "unsupported_formula" for w in result.warnings)
    assert result.omitted_locations == ["Sheet!A1"]
    assert not any("source indexed" in w["message"] for w in result.warnings)
