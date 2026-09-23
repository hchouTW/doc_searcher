"""Stable parse outcomes across formats and edge cases (Task 2.3)."""

import codecs
import os
import sys
from pathlib import Path

import pytest

from doc_searcher.parsers import ParseStatus, get_parser, parse_file

FIXTURES = Path(__file__).resolve().parents[2] / "fixtures" / "files"
SAMPLES = Path(__file__).resolve().parents[2] / "sample_files"
POSIX_NON_ROOT = sys.platform != "win32" and os.geteuid() != 0


# ---------------------------------------------------------------- builders
def make_pdf(path, text=None, **save_kwargs):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    doc.save(str(path), **save_kwargs)
    doc.close()
    return path


def make_docx(path, text=None):
    import docx

    document = docx.Document()
    if text:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def make_xlsx(path, text=None):
    import openpyxl

    workbook = openpyxl.Workbook()
    if text:
        workbook.active["A1"] = text
    workbook.save(str(path))
    return path


def make_pptx(path, text=None):
    import pptx

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    if text:
        slide.shapes.title.text = text
    presentation.save(str(path))
    return path


def texts(result):
    return "\n".join(segment.text for segment in result.segments)


# ------------------------------------------------------------------ empty
@pytest.mark.parametrize(
    "name, build",
    [
        ("empty.pdf", make_pdf),
        ("empty.docx", make_docx),
        ("empty.xlsx", make_xlsx),
        ("empty.pptx", make_pptx),
    ],
)
def test_valid_documents_without_text_are_empty(tmp_path, name, build):
    result = parse_file(str(build(tmp_path / name)))
    assert result.status is ParseStatus.EMPTY
    assert result.error is None
    assert result.segments == []


@pytest.mark.parametrize("content", [b"", b"   \n\t  \r\n"])
def test_blank_text_file_is_empty(tmp_path, content):
    path = tmp_path / "blank.txt"
    path.write_bytes(content)
    assert parse_file(str(path)).status is ParseStatus.EMPTY


# --------------------------------------------------------------- encodings
@pytest.mark.parametrize(
    "encoding, bom",
    [
        ("utf-8", codecs.BOM_UTF8),
        ("utf-16-le", codecs.BOM_UTF16_LE),
        ("utf-16-be", codecs.BOM_UTF16_BE),
        ("utf-16-le", b""),
        ("utf-16-be", b""),
        ("utf-32-le", codecs.BOM_UTF32_LE),
    ],
    ids=["utf8-bom", "utf16le-bom", "utf16be-bom", "utf16le-nobom", "utf16be-nobom", "utf32le-bom"],
)
def test_text_encodings_decode_cleanly(tmp_path, encoding, bom):
    text = "Budget report 2026\n預算報告 第二季\n"
    path = tmp_path / "encoded.txt"
    path.write_bytes(bom + text.encode(encoding))

    result = parse_file(str(path))

    assert result.status is ParseStatus.SUCCESS
    assert texts(result) == text.strip()
    assert "﻿" not in texts(result) and "\x00" not in texts(result)


@pytest.mark.parametrize(
    "raw, expected",
    [
        pytest.param("預算".encode("big5"), "預算"),
        pytest.param(
            "预算报告".encode("gbk"),
            "预算报告",
            marks=pytest.mark.xfail(
                strict=True,
                reason="Known limitation: BOM-less GBK is often also valid Big5, and Big5 "
                "is tried first (Traditional Chinese default)",
            ),
        ),
    ],
    ids=["big5", "gbk"],
)
def test_legacy_cjk_encodings(tmp_path, raw, expected):
    path = tmp_path / "legacy.txt"
    path.write_bytes(raw)
    assert texts(parse_file(str(path))) == expected


def test_invalid_bytes_do_not_fail(tmp_path):
    path = tmp_path / "binary.csv"
    path.write_bytes(b"ok,\xff\xfe\xfa\x80\x81 value")
    result = parse_file(str(path))
    assert result.status is ParseStatus.SUCCESS
    assert "ok" in texts(result) and "value" in texts(result)


# ------------------------------------------------------ corrupt / encrypted
@pytest.mark.parametrize(
    "name", ["bad.pdf", "bad.docx", "bad.xlsx", "bad.pptx", "bad.xls", "bad.ppt", "bad.doc"]
)
@pytest.mark.parametrize(
    "content", [b"", b"\x00\x01garbage not a document" * 40], ids=["zero-bytes", "garbage"]
)
def test_corrupt_documents_are_reported_not_raised(tmp_path, name, content):
    path = tmp_path / name
    path.write_bytes(content)
    result = parse_file(str(path))
    if name == "bad.doc" and content:
        # .doc falls back to salvaging readable strings from non-OLE files (e.g. RTF/text).
        assert result.status in (ParseStatus.SUCCESS, ParseStatus.CORRUPT)
    else:
        assert result.status is ParseStatus.CORRUPT, result.error
        assert result.error


def test_user_password_pdf_is_encrypted(tmp_path):
    import pymupdf

    path = make_pdf(
        tmp_path / "locked.pdf",
        "secret budget",
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        user_pw="user",
        owner_pw="owner",
    )
    result = parse_file(str(path))
    assert result.status is ParseStatus.ENCRYPTED
    assert "password" in result.error


def test_owner_password_only_pdf_is_still_readable(tmp_path):
    import pymupdf

    path = make_pdf(
        tmp_path / "restricted.pdf",
        "restricted budget",
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        owner_pw="owner",
        permissions=pymupdf.PDF_PERM_PRINT,
    )
    result = parse_file(str(path))
    assert result.status is ParseStatus.SUCCESS
    assert "restricted budget" in texts(result)


# ----------------------------------------------------- unreadable / other
def test_missing_file_is_unreadable(tmp_path):
    for name in ("missing.pdf", "missing.docx", "missing.txt"):
        assert parse_file(str(tmp_path / name)).status is ParseStatus.UNREADABLE


@pytest.mark.skipif(not POSIX_NON_ROOT, reason="needs POSIX permissions as non-root")
@pytest.mark.parametrize("name, build", [("locked.docx", make_docx), ("locked.pdf", make_pdf)])
def test_permission_denied_is_unreadable(tmp_path, name, build):
    path = build(tmp_path / name, "text")
    os.chmod(path, 0)
    try:
        result = parse_file(str(path))
    finally:
        os.chmod(path, 0o644)
    assert result.status is ParseStatus.UNREADABLE


def test_unsupported_extension(tmp_path):
    path = tmp_path / "photo.png"
    path.write_bytes(b"\x89PNG")
    assert parse_file(str(path)).status is ParseStatus.UNSUPPORTED


@pytest.mark.parametrize(
    "module, name, build",
    [
        ("pymupdf", "a.pdf", make_pdf),
        ("docx", "a.docx", make_docx),
        ("openpyxl", "a.xlsx", make_xlsx),
        ("pptx", "a.pptx", make_pptx),
    ],
)
def test_missing_library_is_dependency_missing(tmp_path, monkeypatch, module, name, build):
    path = build(tmp_path / name, "text")
    monkeypatch.setitem(sys.modules, module, None)
    monkeypatch.setitem(sys.modules, "fitz", None)
    result = parse_file(str(path))
    assert result.status is ParseStatus.DEPENDENCY_MISSING
    assert "not installed" in result.error


def test_parser_crash_becomes_corrupt_result(tmp_path, monkeypatch):
    path = make_docx(tmp_path / "a.docx", "text")
    monkeypatch.setattr(get_parser(str(path)), "parse", lambda p: 1 / 0)
    result = parse_file(str(path))
    assert result.status is ParseStatus.CORRUPT
    assert "Parser exception" in result.error


# --------------------------------------------------------------- paths
@pytest.mark.parametrize(
    "name",
    [
        "報告 最終版 (1).pdf",
        "Résumé 2026.docx",
        "予算 表.xlsx",
        "간단한 발표.pptx",
        "notes with spaces.txt",
    ],
)
def test_non_ascii_and_space_paths(tmp_path, name):
    folder = tmp_path / "資料 夾 ☃"
    folder.mkdir()
    path = folder / name
    builders = {".pdf": make_pdf, ".docx": make_docx, ".xlsx": make_xlsx, ".pptx": make_pptx}
    suffix = Path(name).suffix
    if suffix in builders:
        builders[suffix](path, "Budget 2026")
    else:
        path.write_text("Budget 2026", encoding="utf-8")
    result = parse_file(str(path))
    assert result.status is ParseStatus.SUCCESS, result.error
    assert "Budget 2026" in texts(result)


# ------------------------------------------------------------ legacy formats
def test_real_word97_doc_is_extracted_best_effort():
    result = parse_file(str(FIXTURES / "legacy_word97.doc"))
    assert result.status is ParseStatus.SUCCESS
    assert "Legacy Word fixture" in texts(result)
    assert "舊版文件預算 test budget" in texts(result)


def test_real_xls_releases_its_file_handle(monkeypatch):
    import xlrd

    released = []
    real_open = xlrd.open_workbook

    def tracking_open(*args, **kwargs):
        book = real_open(*args, **kwargs)
        real_release = book.release_resources
        book.release_resources = lambda: (released.append(True), real_release())
        return book

    monkeypatch.setattr(xlrd, "open_workbook", tracking_open)
    result = parse_file(str(SAMPLES / "sample_legacy_financial.xls"))
    assert result.status is ParseStatus.SUCCESS
    assert released == [True]


def test_pdf_document_is_closed(tmp_path, monkeypatch):
    import pymupdf

    closed = []
    real_close = pymupdf.Document.close
    monkeypatch.setattr(
        pymupdf.Document, "close", lambda self: (closed.append(True), real_close(self))
    )
    parse_file(str(make_pdf(tmp_path / "a.pdf", "text")))
    parse_file(
        str(
            make_pdf(
                tmp_path / "b.pdf",
                "x",
                encryption=pymupdf.PDF_ENCRYPT_AES_256,
                user_pw="u",
                owner_pw="o",
            )
        )
    )
    assert len(closed) >= 2
