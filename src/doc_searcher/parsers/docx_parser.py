# Purpose: Modern Word (.docx) document text extraction parser.
# What the code does:
#   - Uses python-docx to extract paragraphs, tables, and structural text.
#   - Extracts table cell contents in reading order.
#   - Handles protected or malformed files gracefully without unhandled exceptions.
# Usage notes, dependencies, or assumptions:
#   - Requires python-docx.
#   - Returns PageSegment grouped by structural sections or blocks.

import os
from typing import List
from .base import BaseParser, ExtractedDoc, PageSegment, ParseStatus


class DocxParser(BaseParser):
    """Parser for Microsoft Word (.docx) documents."""

    def parse(self, file_path: str) -> ExtractedDoc:
        abs_path = os.path.abspath(file_path)
        segments: List[PageSegment] = []

        try:
            import docx
        except ImportError:
            return ExtractedDoc.failed(
                abs_path, "docx", ParseStatus.DEPENDENCY_MISSING, "python-docx is not installed."
            )

        try:
            doc = docx.Document(abs_path)
        except Exception as e:
            return ExtractedDoc.from_exception(abs_path, "docx", "Cannot open docx file", e)

        try:
            # Collect paragraphs
            para_texts = []
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    para_texts.append(text)

            # Collect tables
            table_texts = []
            for table_idx, table in enumerate(doc.tables, start=1):
                row_strings = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    # Deduplicate repeated cells from merged columns
                    unique_cells = []
                    for c in row_data:
                        if not unique_cells or unique_cells[-1] != c:
                            unique_cells.append(c)
                    if unique_cells:
                        row_strings.append(" | ".join(unique_cells))
                if row_strings:
                    table_texts.append(f"[Table {table_idx}]\n" + "\n".join(row_strings))

            # Group content into coherent segments
            seg_idx = 1
            if para_texts:
                segments.append(
                    PageSegment(
                        segment_id=f"Section {seg_idx} (Paragraphs)",
                        segment_type="section",
                        text="\n\n".join(para_texts)
                    )
                )
                seg_idx += 1

            if table_texts:
                segments.append(
                    PageSegment(
                        segment_id=f"Section {seg_idx} (Tables)",
                        segment_type="section",
                        text="\n\n".join(table_texts)
                    )
                )

            return ExtractedDoc(
                file_path=abs_path,
                file_type="docx",
                total_segments=len(segments),
                segments=segments
            )
        except Exception as e:
            return ExtractedDoc.from_exception(abs_path, "docx", "Error reading docx content", e)
