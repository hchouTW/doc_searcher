# Purpose: Modern Excel (.xlsx) workbook text extraction parser.
# What the code does:
#   - Uses openpyxl with read_only=True and data_only=True for low-memory streaming.
#   - Iterates through worksheets and non-empty rows.
#   - Formats rows as readable tabular text segments keyed by sheet name.
# Usage notes, dependencies, or assumptions:
#   - Requires openpyxl.
#   - Returns PageSegment per worksheet.

import os
from typing import List
from .base import BaseParser, ExtractedDoc, PageSegment, ParseStatus


class XlsxParser(BaseParser):
    """Parser for Microsoft Excel (.xlsx) spreadsheets using openpyxl."""

    def parse(self, file_path: str) -> ExtractedDoc:
        abs_path = os.path.abspath(file_path)
        segments: List[PageSegment] = []

        try:
            import openpyxl
        except ImportError:
            return ExtractedDoc.failed(
                abs_path, "xlsx", ParseStatus.DEPENDENCY_MISSING, "openpyxl is not installed."
            )

        wb = None
        try:
            wb = openpyxl.load_workbook(abs_path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames

            for sheet_name in sheet_names:
                sheet = wb[sheet_name]
                row_lines = []

                for row in sheet.iter_rows(values_only=True):
                    # Filter non-None cells
                    row_vals = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                    if row_vals:
                        row_lines.append(" | ".join(row_vals))

                sheet_text = "\n".join(row_lines).strip()
                if sheet_text:
                    segments.append(
                        PageSegment(
                            segment_id=sheet_name,
                            segment_type="sheet",
                            text=sheet_text
                        )
                    )

            return ExtractedDoc(
                file_path=abs_path,
                file_type="xlsx",
                total_segments=len(sheet_names),
                segments=segments
            )
        except Exception as e:
            return ExtractedDoc.from_exception(abs_path, "xlsx", "Error reading xlsx file", e)
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass
